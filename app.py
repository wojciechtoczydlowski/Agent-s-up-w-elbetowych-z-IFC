import os
import ifcopenshell
import ifcopenshell.geom
import re
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
from openpyxl.styles import Font
import streamlit as st

from langchain_core.tools import tool
from langchain_openai import ChatOpenAI

# BEZPIECZEŃSTWO (Lokalnie + Chmura): 
# 1. Sprawdzamy czy klucz jest w sekretach Streamlita (gdy aplikacja jest w chmurze)
try:
    if "OPENAI_API_KEY" in st.secrets:
        os.environ["OPENAI_API_KEY"] = st.secrets["OPENAI_API_KEY"]
except Exception:
    pass

# 2. Jeśli nadal nie ma klucza, a uruchamiasz lokalnie, wpisz go tutaj (ale pamiętaj, by go usunąć przed pushem na GitHub!):
if "OPENAI_API_KEY" not in os.environ or not os.environ["OPENAI_API_KEY"]:
    # Wklej tu swój nowy klucz tylko do testów lokalnych:
    os.environ["OPENAI_API_KEY"] = ""

@tool
def generuj_zestawienie_slupow(file_path: str) -> str:
    """
    Narzędzie służące do precyzyjnej analizy słupów żelbetowych w pliku IFC. 
    Wczytuje model, znajduje słupy, wykrywa wsporniki, oblicza wysokość, 
    koryguje błędy wymiarowe, grupuje je i generuje plik Excel (.xlsx).
    """
    try:
        if not os.path.exists(file_path):
            return f"Błąd: Nie znaleziono pliku {file_path}."

        model = ifcopenshell.open(file_path)
        columns = model.by_type("IfcColumn")
        
        settings = ifcopenshell.geom.settings()
        settings.set(settings.USE_WORLD_COORDS, True)

        zestawienie_dokladne = {}
        bledy = 0

        for col in columns: 
            nazwa_surowa = col.Name if col.Name else ""
            if not nazwa_surowa or "concrete" not in nazwa_surowa.lower():
                if col.ObjectType:
                    nazwa_surowa = col.ObjectType

            if nazwa_surowa and "concrete" in nazwa_surowa.lower():
                dopasowanie = re.search(r'(\d+)x(\d+)', nazwa_surowa)
                uwaga_do_elementu = ""
                
                if dopasowanie:
                    wymiar_1, wymiar_2 = int(dopasowanie.group(1)), int(dopasowanie.group(2))
                    oryginalny_przekroj = f"{wymiar_1}x{wymiar_2}"
                    skorygowano = False
                    if wymiar_1 < 150: wymiar_1 *= 10; skorygowano = True
                    if wymiar_2 < 150: wymiar_2 *= 10; skorygowano = True
                    przekroj = f"{wymiar_1}x{wymiar_2}"
                    if skorygowano: uwaga_do_elementu = f"Skorygowano błąd wymiaru: z {oryginalny_przekroj} na {przekroj}"
                    nazwa_bazowa = nazwa_surowa.split(oryginalny_przekroj)[0].strip(':').strip()
                else:
                    przekroj = "Brak"
                    czesc = nazwa_surowa.split(':')
                    if len(czesc) > 1 and czesc[-1].isdigit():
                        nazwa_bazowa = ":".join(czesc[:-1]).strip()
                    else:
                        nazwa_bazowa = nazwa_surowa
                
                # === POPRAWKA: IGNOROWANIE PRZYROSTKÓW "MY", "MY 2" itp. ===
                # Wyłapuje i usuwa końcówki z wielkimi literami MY oraz ewentualnymi cyframi
                nazwa_bazowa = re.sub(r'\s*MY\s*\d*', '', nazwa_bazowa).strip()
                # ============================================================
                
                try:
                    shape = ifcopenshell.geom.create_shape(settings, col)
                    verts = shape.geometry.verts
                    
                    z_coords = [verts[j] for j in range(2, len(verts), 3)]
                    z_min = min(z_coords)
                    z_max = max(z_coords)
                    
                    height = (z_max - z_min) + 0.05
                    height_rounded = round(height, 2)
                    
                    ma_wspornik = any((z_min + 0.3) < z < (z_max - 0.3) for z in z_coords)
                    typ_geometrii = "Ze wspornikiem" if ma_wspornik else "Zwykły"
                    
                    klucz_slupa = (nazwa_bazowa, przekroj, typ_geometrii, height_rounded)
                    
                    if klucz_slupa not in zestawienie_dokladne:
                        zestawienie_dokladne[klucz_slupa] = {'ilosc': 0, 'uwagi': set()}
                        
                    zestawienie_dokladne[klucz_slupa]['ilosc'] += 1
                    if uwaga_do_elementu:
                        zestawienie_dokladne[klucz_slupa]['uwagi'].add(uwaga_do_elementu)
                except:
                    bledy += 1

        zestawienie_zaokraglone = {}
        suma_calkowita = 0
        for klucz, dane in zestawienie_dokladne.items():
            nazwa, przekroj, typ_bryly, wysokosc = klucz
            wys_zaokr = float(Decimal(str(wysokosc)).quantize(Decimal('0.1'), rounding=ROUND_HALF_UP))
            nowy_klucz = (nazwa, przekroj, typ_bryly, wys_zaokr)
            
            if nowy_klucz not in zestawienie_zaokraglone:
                zestawienie_zaokraglone[nowy_klucz] = {'ilosc': 0, 'uwagi': set()}
                
            zestawienie_zaokraglone[nowy_klucz]['ilosc'] += dane['ilosc']
            zestawienie_zaokraglone[nowy_klucz]['uwagi'].update(dane['uwagi'])
            suma_calkowita += dane['ilosc']

        def zasady_sortowania(element):
            klucz, _ = element
            _, przekroj, _, _ = klucz
            try:
                w1, w2 = map(int, przekroj.split('x'))
            except:
                w1, w2 = 0, 0
            return (w1, w2, -klucz[3], klucz[2])

        baza_pliku = os.path.splitext(os.path.basename(file_path))[0]
        nazwa_pliku = f"zestawienie_slupow_{baza_pliku}.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Zestawienie Słupów"
        
        def dodaj_wiersz(d, pogrubiony=False):
            ws.append(d)
            if pogrubiony:
                for col in range(1, len(d) + 1):
                    ws.cell(row=ws.max_row, column=col).font = Font(bold=True)

        dodaj_wiersz(['--- ZESTAWIENIE SŁUPÓW ŻELBETOWYCH ---'], pogrubiony=True)
        dodaj_wiersz(['Nazwa bazy', 'Przekrój', 'Typ bryły', 'Wysokość zaokrąglona (m)', 'Ilość sztuk', 'Uwagi'], pogrubiony=True)

        for klucz, dane in sorted(zestawienie_zaokraglone.items(), key=zasady_sortowania):
            dodaj_wiersz([klucz[0], klucz[1], klucz[2], klucz[3], dane['ilosc'], " | ".join(dane['uwagi'])])
            
        dodaj_wiersz(['', '', '', 'SUMA SŁUPÓW:', suma_calkowita, ''], pogrubiony=True)

        for kolumna in ws.columns:
            max_dlugosc = max((len(str(komorka.value)) for komorka in kolumna if komorka.value), default=0)
            ws.column_dimensions[kolumna[0].column_letter].width = max_dlugosc + 2

        wb.save(nazwa_pliku)
        return f"Sukces! Przeanalizowano model '{baza_pliku}'. Znaleziono łącznie {suma_calkowita} słupów. Zapisano raport."

    except Exception as e:
        return f"Wystąpił błąd podczas analizy pliku IFC: {str(e)}"

# ==========================================
# INTERFEJS STREAMLIT
# ==========================================
st.set_page_config(page_title="LLENTAB BIM Agent", page_icon="🏗️", layout="wide")
st.title("🏗️ Asystent BIM: Dynamiczny Przedmiar Słupów")

st.sidebar.header("Wgraj model IFC")
wgrany_plik = st.sidebar.file_uploader("Wybierz plik .ifc", type=["ifc"])

if wgrany_plik is not None:
    sciezka_tymczasowa = wgrany_plik.name
    with open(sciezka_tymczasowa, "wb") as f:
        f.write(wgrany_plik.getbuffer())
    st.sidebar.success(f"Wczytano plik: {wgrany_plik.name}")
else:
    sciezka_tymczasowa = "PL7258_I.ifc"

if st.sidebar.button("🚀 Uruchom analizę modelu"):
    with st.spinner("Przetwarzanie geometrii..."):
        wynik = generuj_zestawienie_slupow.invoke({"file_path": sciezka_tymczasowa})
        st.session_state["wynik"] = wynik
        st.session_state["aktywny_plik"] = sciezka_tymczasowa
        st.success("Zakończono!")

if "wynik" in st.session_state:
    st.info(st.session_state["wynik"])
    baza_biezaca = os.path.splitext(os.path.basename(st.session_state.get("aktywny_plik", sciezka_tymczasowa)))[0]
    generowany_excel = f"zestawienie_slupow_{baza_biezaca}.xlsx"
    
    if os.path.exists(generowany_excel):
        with open(generowany_excel, "rb") as f:
            st.download_button("📥 Pobierz Excel", f, file_name=generowany_excel)

st.divider()
st.subheader("💬 Czat z modelem")

if "messages" not in st.session_state:
    st.session_state["messages"] = []

for msg in st.session_state["messages"]:
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])

if prompt := st.chat_input("Zadaj pytanie..."):
    st.session_state["messages"].append({"role": "user", "content": prompt})
    with st.chat_message("user"):
        st.markdown(prompt)

    with st.chat_message("assistant"):
        # Ważne, by model miał do dyspozycji Twój klucz API (zapewniony u góry skryptu)
        llm = ChatOpenAI(model="gpt-4o", temperature=0)
        llm_z_narzedziami = llm.bind_tools([generuj_zestawienie_slupow])
        
        history = "\n".join([f"{m['role']}: {m['content']}" for m in st.session_state["messages"]])
        odpowiedz = llm_z_narzedziami.invoke(history)
        
        if odpowiedz.tool_calls:
            for tc in odpowiedz.tool_calls:
                if tc["name"] == "generuj_zestawienie_slupow":
                    args = tc["args"]
                    args["file_path"] = st.session_state.get("aktywny_plik", sciezka_tymczasowa)
                    res = generuj_zestawienie_slupow.invoke(args)
                    st.markdown(res)
                    st.session_state["messages"].append({"role": "assistant", "content": res})
        else:
            st.markdown(odpowiedz.content)
            st.session_state["messages"].append({"role": "assistant", "content": odpowiedz.content})